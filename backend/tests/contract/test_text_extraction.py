"""TextExtractionPort: the common subset over the fake and the real PDF/DOCX extractor, plus real-file rejections."""

import importlib.util
import io
from collections.abc import Callable
from pathlib import Path
from typing import Any

import pytest
from tests.fakes.extraction import FakeTextExtractor

from scout.vault.adapters import extraction as extraction_module
from scout.vault.adapters.extraction import PdfDocxExtractor
from scout.vault.domain.enums import DocumentMime, PageBasis
from scout.vault.ports.extraction import ExtractionRejected, TextExtractionPort

_SPEC = importlib.util.spec_from_file_location(
    "make_sample_docs", Path(__file__).resolve().parents[3] / "scripts" / "make_sample_docs.py"
)
assert _SPEC and _SPEC.loader
samples = importlib.util.module_from_spec(_SPEC)
_SPEC.loader.exec_module(samples)

QUOTE = "The applicants do not demonstrate that the foams meet fire-safety standard EN 13501"
PROPOSAL = (Path(samples.SOURCE) / "declined-proposal-lignin-foams.md").read_text(encoding="utf-8")


def _fake() -> tuple[TextExtractionPort, Callable[[], tuple[bytes, DocumentMime]]]:
    pages = ["Aims. " * 40, "Methods. " * 40, f"Reviewer 2: {QUOTE} at the proposed bio-content."]
    return FakeTextExtractor(), lambda: ("\f".join(pages).encode(), DocumentMime.PDF)


def _pdf() -> tuple[TextExtractionPort, Callable[[], tuple[bytes, DocumentMime]]]:
    return PdfDocxExtractor(), lambda: (samples.to_pdf(PROPOSAL), DocumentMime.PDF)


def _docx() -> tuple[TextExtractionPort, Callable[[], tuple[bytes, DocumentMime]]]:
    return PdfDocxExtractor(), lambda: (samples.to_docx(PROPOSAL), DocumentMime.DOCX)


@pytest.fixture(
    params=[pytest.param(_fake, id="fake"), pytest.param(_pdf, id="pdf"), pytest.param(_docx, id="docx")]
)
def case(
    request: pytest.FixtureRequest,
) -> tuple[TextExtractionPort, Callable[[], tuple[bytes, DocumentMime]]]:
    return request.param()  # type: ignore[no-any-return]


async def test_pages_and_quote_on_the_right_page(
    case: tuple[TextExtractionPort, Callable[[], tuple[bytes, DocumentMime]]],
) -> None:
    extractor, data = case
    result = await extractor.extract_text(*data())
    assert len(result.pages) == 3
    assert QUOTE in result.pages[2].replace("\n", " ") and QUOTE not in result.pages[0]
    assert result.extractor


# --- real extractor ------------------------------------------------------------------------------


def _pdf_with(draw: Callable[[Any], None], pages: int = 1, encrypt: str | None = None) -> bytes:
    from reportlab.lib import pdfencrypt
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, encrypt=pdfencrypt.StandardEncryption(encrypt, canPrint=0) if encrypt else None)
    for _ in range(pages):
        draw(c)
        c.showPage()
    c.save()
    return buf.getvalue()


async def test_pdf_page_basis_and_docx_basis() -> None:
    e = PdfDocxExtractor()
    pdf = await e.extract_text(samples.to_pdf(PROPOSAL), DocumentMime.PDF)
    docx = await e.extract_text(samples.to_docx(PROPOSAL), DocumentMime.DOCX)
    assert (pdf.page_basis, pdf.extractor) == (PageBasis.PHYSICAL, "pypdf")
    assert (docx.page_basis, docx.extractor) == (PageBasis.PAGE_BREAKS, "python-docx")


async def test_docx_without_page_breaks_has_no_page_basis_and_keeps_tables() -> None:
    import docx

    d = docx.Document()
    d.add_paragraph("Intro paragraph. " * 20)
    t = d.add_table(rows=2, cols=2)
    for i, cell in enumerate(t.rows[0].cells + t.rows[1].cells):
        cell.text = f"cell{i}"
    buf = io.BytesIO()
    d.save(buf)
    r = await PdfDocxExtractor().extract_text(buf.getvalue(), DocumentMime.DOCX)
    assert r.page_basis is PageBasis.NONE and len(r.pages) == 1
    assert "cell0 | cell1" in r.pages[0] and "cell2 | cell3" in r.pages[0]


async def test_encrypted_pdf_is_rejected() -> None:
    data = _pdf_with(lambda c: c.drawString(72, 700, "secret " * 50), encrypt="owner-and-user-pw")
    with pytest.raises(ExtractionRejected) as e:
        await PdfDocxExtractor().extract_text(data, DocumentMime.PDF)
    assert e.value.reject_code == "encrypted"


async def test_scanned_pdf_has_no_text_layer() -> None:
    data = _pdf_with(lambda c: c.rect(50, 50, 400, 600, fill=1), pages=3)  # drawings only, like a scan
    with pytest.raises(ExtractionRejected) as e:
        await PdfDocxExtractor().extract_text(data, DocumentMime.PDF)
    assert e.value.reject_code == "no_text_layer"


async def test_too_many_pages_and_too_long() -> None:
    data = _pdf_with(lambda c: c.drawString(72, 700, "word " * 60), pages=5)
    with pytest.raises(ExtractionRejected) as e:
        await PdfDocxExtractor(max_pages=4).extract_text(data, DocumentMime.PDF)
    assert e.value.reject_code == "too_many_pages"
    with pytest.raises(ExtractionRejected) as e2:
        await PdfDocxExtractor(max_chars=500).extract_text(data, DocumentMime.PDF)
    assert e2.value.reject_code == "text_too_long"


@pytest.mark.parametrize("mime", [DocumentMime.PDF, DocumentMime.DOCX])
async def test_corrupt_files(mime: DocumentMime) -> None:
    with pytest.raises(ExtractionRejected) as e:
        await PdfDocxExtractor().extract_text(b"%PDF-1.7 this is not a pdf" + b"\x00" * 50, mime)
    assert e.value.reject_code == "corrupt_file"


async def test_timeout_kills_the_child(monkeypatch: pytest.MonkeyPatch) -> None:
    async def never(coro: Any, timeout: float) -> Any:  # noqa: ASYNC109
        coro.close()
        raise TimeoutError

    monkeypatch.setattr(extraction_module.asyncio, "wait_for", never)
    with pytest.raises(ExtractionRejected) as e:
        await PdfDocxExtractor().extract_text(samples.to_pdf(PROPOSAL), DocumentMime.PDF)
    assert e.value.reject_code == "extraction_timeout"


async def test_thirty_page_pdf_is_fast_and_complete() -> None:
    data = samples.to_pdf(samples.long_proposal())
    r = await PdfDocxExtractor().extract_text(data, DocumentMime.PDF)
    assert len(r.pages) == 30 and QUOTE in " ".join(r.pages[28].split())
