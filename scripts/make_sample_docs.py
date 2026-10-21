#!/usr/bin/env python3
"""Render the fictional sample vault documents (``tests/fixtures/vault/*.md``) to PDF and DOCX.

Output goes to ``tests/fixtures/vault/build/`` (git-ignored). ``<!-- page -->`` in the Markdown marks a page
break. ``--long`` also writes a 30-page PDF declined proposal (for the "ready within 2 minutes" acceptance test).

Usage: ``cd backend && uv run python ../scripts/make_sample_docs.py [--long] [--out DIR]``
"""

from __future__ import annotations

import argparse
import io
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "tests" / "fixtures" / "vault"
PAGE = "<!-- page -->"
FILLER = (
    "Lignin is a heterogeneous aromatic biopolymer. Its hydroxyl content depends on the pulping process and on "
    "the fractionation route, which in turn influences reactivity towards isocyanates in foam formulations. "
)


def _inline(text: str) -> str:
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    return re.sub(r"`(.+?)`", r"<font face='Courier'>\1</font>", text)


def blocks(markdown: str) -> list[tuple[str, str]]:
    """Split Markdown into ``(kind, text)`` blocks: h1, h2, p, quote, page."""
    out: list[tuple[str, str]] = []
    para: list[str] = []

    def flush() -> None:
        if para:
            out.append(("p", " ".join(para)))
            para.clear()

    for line in markdown.splitlines():
        s = line.strip()
        if not s:
            flush()
        elif s == PAGE:
            flush()
            out.append(("page", ""))
        elif s.startswith("# "):
            flush()
            out.append(("h1", s[2:]))
        elif s.startswith("## "):
            flush()
            out.append(("h2", s[3:]))
        elif s.startswith("> "):
            flush()
            out.append(("quote", s[2:]))
        elif re.match(r"(\d+\.|-)\s", s):
            flush()
            out.append(("p", s))
        else:
            para.append(s)
    flush()
    return out


def to_pdf(markdown: str) -> bytes:
    """Render with reportlab (text layer, physical pages)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    flow: list[object] = []
    for kind, text in blocks(markdown):
        if kind == "page":
            flow.append(PageBreak())
        else:
            style = {"h1": "Heading1", "h2": "Heading2", "quote": "Italic"}.get(kind, "BodyText")
            flow += [Paragraph(_inline(text), styles[style]), Spacer(1, 6)]
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, title="Sample document").build(flow)
    return buf.getvalue()


def to_docx(markdown: str) -> bytes:
    """Render with python-docx (explicit page breaks)."""
    import docx

    d = docx.Document()
    for kind, text in blocks(markdown):
        if kind == "page":
            d.add_page_break()
        elif kind == "h1":
            d.add_heading(text, level=1)
        elif kind == "h2":
            d.add_heading(text, level=2)
        else:
            d.add_paragraph(re.sub(r"[*`]", "", text))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def long_proposal(pages: int = 30) -> str:
    """A fictional declined proposal of ``pages`` pages; reviewer comments sit on the second to last page."""
    parts = ["> SAMPLE DOCUMENT — fictional, written for Lattice Scout testing.", "# Long sample proposal (declined)"]
    for n in range(1, pages - 1):
        parts += [f"## Work package {n}", FILLER * 6, FILLER * 4, PAGE]
    parts += [
        "## Panel summary and reviewer comments",
        "Reviewer 2: The applicants do not demonstrate that the foams meet fire-safety standard EN 13501 at the "
        "proposed bio-content.",
        "Panel: Strong science, but the application case is not yet convincing.",
        PAGE,
        "## Budget",
        "The requested budget is detailed in the attached form.",
    ]
    return "\n\n".join(parts)


def main() -> None:
    """Write the files."""
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--out", type=Path, default=SOURCE / "build")
    ap.add_argument("--long", action="store_true", help="also write sample-long-proposal-30p.pdf")
    args = ap.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    for src in sorted(SOURCE.glob("*.md")):
        if src.name == "README.md":
            continue
        md = src.read_text(encoding="utf-8")
        (args.out / f"{src.stem}.pdf").write_bytes(to_pdf(md))
        (args.out / f"{src.stem}.docx").write_bytes(to_docx(md))
        print(f"wrote {src.stem}.pdf / .docx")
    if args.long:
        (args.out / "sample-long-proposal-30p.pdf").write_bytes(to_pdf(long_proposal()))
        print("wrote sample-long-proposal-30p.pdf")


if __name__ == "__main__":
    main()
